import openpyxl
from collections import Counter


def analyze_survey_detail(filepath, name):
    output = []
    output.append(f"\n{'=' * 70}")
    output.append(f"DETAILED SURVEY ANALYSIS: {name}")
    output.append(f"{'=' * 70}")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.worksheets[0]

    output.append(f"\nTotal responses: {ws.max_row - 1}")  # minus header

    # Analyze each question column
    for col in range(2, ws.max_column + 1):
        header = ws.cell(1, col).value
        if not header:
            continue

        # Get all responses for this question
        responses = []
        for row in range(2, ws.max_row + 1):
            val = ws.cell(row, col).value
            if val:
                responses.append(str(val).strip())

        unique_responses = Counter(responses)
        output.append(f"\n--- Question {col - 1} ---")
        output.append(f"Header: {header}")
        output.append(f"Total responses: {len(responses)}")
        output.append(f"Unique values: {len(unique_responses)}")

        if unique_responses:
            output.append("Response distribution:")
            for resp, count in unique_responses.most_common(10):
                output.append(f"  {resp[:60]}: {count}")

    return "\n".join(output)


def analyze_template_detail(filepath, name):
    output = []
    output.append(f"\n{'=' * 70}")
    output.append(f"DETAILED TEMPLATE ANALYSIS: {name}")
    output.append(f"{'=' * 70}")

    wb = openpyxl.load_workbook(filepath, data_only=True)

    if "Tabulación" in wb.sheetnames:
        ws = wb["Tabulación"]
        output.append(f"\n--- Tabulación Sheet ({ws.max_row} rows) ---")

        # Find all question blocks
        for row in range(1, ws.max_row + 1):
            cell_val = ws.cell(row, 2).value
            if cell_val and isinstance(cell_val, str) and "Cuadro" in cell_val:
                output.append(f"Row {row}: {cell_val[:80]}")

                # Show next few rows
                for r in range(row + 1, min(row + 15, ws.max_row + 1)):
                    row_data = [ws.cell(r, c).value for c in range(1, 7)]
                    if any(row_data):
                        output.append(f"  Row {r}: {row_data[:4]}")

    return "\n".join(output)


# Run analyses
files = [
    (
        r"D:\Tareas\Año_2026\3.Marzo\plan_next\plantillas\excel\fase 2\Encuesta.xlsx",
        "Survey",
    ),
    (
        r"D:\Tareas\Año_2026\3.Marzo\plan_next\plantillas\excel\fase 2\plantilla 2.xlsx",
        "Template",
    ),
    (
        r"D:\Tareas\Año_2026\3.Marzo\plan_next\plantillas\excel\fase 2\tabulacion_1773777854.xlsx",
        "Tabulacion",
    ),
]

all_output = []

for filepath, name in files:
    try:
        if name == "Survey":
            result = analyze_survey_detail(filepath, name)
        else:
            result = analyze_template_detail(filepath, name)
        all_output.append(result)
    except Exception as e:
        all_output.append(f"Error: {e}")
        import traceback

        all_output.append(traceback.format_exc())

# Write to file
with open(
    r"D:\Tareas\Año_2026\3.Marzo\plan_next\analysis_detail.txt", "w", encoding="utf-8"
) as f:
    f.write("\n".join(all_output))

print("Detailed analysis written to analysis_detail.txt")
