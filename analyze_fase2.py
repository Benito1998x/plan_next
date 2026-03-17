import openpyxl
import sys

# Set UTF-8 output
sys.stdout.reconfigure(encoding="utf-8")


def analyze_excel(filepath, name):
    output = []
    output.append(f"\n{'=' * 60}")
    output.append(f"FILE: {name}")
    output.append(f"Path: {filepath}")
    output.append(f"{'=' * 60}")

    wb = openpyxl.load_workbook(filepath, data_only=True)

    output.append(f"\nSheet names: {wb.sheetnames}")
    output.append(f"Total sheets: {len(wb.sheetnames)}")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        output.append(f"\n--- Sheet: '{sheet_name}' ---")
        output.append(f"Rows: {ws.max_row}, Columns: {ws.max_column}")

        # Get headers (first row)
        headers = []
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(1, col).value
            headers.append(cell_value)

        output.append(f"\nColumn headers ({len(headers)} columns):")
        for i, h in enumerate(headers, 1):
            output.append(f"  {i}. {h}")

        # Show first few data rows
        if ws.max_row > 1:
            output.append(f"\nFirst 3 data rows (showing first 10 columns):")
            for row_idx in range(2, min(5, ws.max_row + 1)):
                row_data = []
                for col in range(1, min(11, ws.max_column + 1)):
                    val = ws.cell(row_idx, col).value
                    row_data.append(str(val)[:30] if val is not None else "")
                output.append(f"  Row {row_idx}: {row_data}")

        # Look for question/response type indicators
        output.append(f"\nQuestions found:")
        for row_idx in range(1, min(ws.max_row + 1, 60)):
            for col_idx in range(1, min(ws.max_column + 1, 10)):
                val = ws.cell(row_idx, col_idx).value
                if val and isinstance(val, str):
                    lower_val = val.lower()
                    if "cuadro" in lower_val and "pregunta" in lower_val:
                        output.append(f"  Row {row_idx}, Col {col_idx}: {val[:100]}")
                        break

    return "\n".join(output)


# Main analysis
files = [
    (
        r"D:\Tareas\Año_2026\3.Marzo\plan_next\plantillas\excel\fase 2\Encuesta.xlsx",
        "Encuesta (Survey)",
    ),
    (
        r"D:\Tareas\Año_2026\3.Marzo\plan_next\plantillas\excel\fase 2\plantilla 2.xlsx",
        "Plantilla 2 (Template v2)",
    ),
    (
        r"D:\Tareas\Año_2026\3.Marzo\plan_next\plantillas\excel\fase 2\tabulacion_1773777854.xlsx",
        "Tabulacion (Tabulation)",
    ),
]

all_output = []

for filepath, name in files:
    try:
        result = analyze_excel(filepath, name)
        all_output.append(result)
    except Exception as e:
        all_output.append(f"\nError analyzing {name}: {e}")
        import traceback

        all_output.append(traceback.format_exc())

# Write to file
with open(
    r"D:\Tareas\Año_2026\3.Marzo\plan_next\analysis_output.txt", "w", encoding="utf-8"
) as f:
    f.write("\n".join(all_output))

print("Analysis complete. Output written to analysis_output.txt")
