"""
Word Service - Genera documentos Word desde datos de Excel.

Usa python-docx para crear documentos con tablas y formato.
"""

from pathlib import Path
from typing import Dict, Any, List, Optional
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from core.exceptions import ProcessingError, ErrorCodes


class WordService:
    """
    Servicio para generar documentos Word desde datos de Excel.

    Crea documentos con:
    - Título del plan de negocio
    - Sección de Parámetros Globales (tabla)
    - Sección de Productos/Servicios (tabla)

    Usage:
        service = WordService()
        service.generate_from_excel_data(
            {"nombre": "Shawarma Cruz", ...},
            [{"nombre": "Shawarma Cordero", ...}],
            output_path
        )
    """

    def __init__(self):
        """Inicializa el servicio Word."""
        pass

    def generate_from_excel_data(
        self,
        parametros: Dict[str, Any],
        productos: List[Dict[str, Any]],
        output_path: Path,
        title: Optional[str] = None,
    ) -> Path:
        """
        Genera documento Word desde datos extraídos de Excel.

        Args:
            parametros: Dict con parámetros globales
            productos: Lista de productos
            output_path: Ruta donde guardar el documento
            title: Título del documento (opcional, usa nombre del proyecto)

        Returns:
            Path al documento generado

        Raises:
            ProcessingError: Si hay error al generar
        """
        try:
            doc = Document()

            # ===== TÍTULO =====
            nombre_proyecto = parametros.get("nombre", "Plan de Negocio")
            title = title or f"PLAN DE NEGOCIO: {nombre_proyecto}"

            title_paragraph = doc.add_heading(title, level=0)
            title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Espacio
            doc.add_paragraph()

            # ===== SECCIÓN 1: PARÁMETROS GLOBALES =====
            self._add_section_title(doc, "1. Datos Globales")
            self._add_parameters_table(doc, parametros)

            # Espacio
            doc.add_paragraph()

            # ===== SECCIÓN 2: PRODUCTOS / SERVICIOS =====
            self._add_section_title(doc, "2. Productos / Servicios")
            self._add_products_table(doc, productos)

            # Guardar documento
            doc.save(output_path)

            return output_path

        except Exception as e:
            raise ProcessingError(
                message=f"Error al generar documento Word: {str(e)}",
                file_name=str(output_path),
                operation="generar",
                code=ErrorCodes.WORD_GENERATION_ERROR,
            )

    def generate_from_plan(self, plan_data: Dict[str, Any], output_path: Path) -> Path:
        """
        Genera documento Word desde un plan completo.

        Args:
            plan_data: Dict con plan completo:
                {
                    "parametros_globales": {...},
                    "productos": [...]
                }
            output_path: Ruta donde guardar

        Returns:
            Path al documento generado
        """
        parametros = plan_data.get("parametros_globales", {})
        productos = plan_data.get("productos", [])

        return self.generate_from_excel_data(parametros, productos, output_path)

    def _add_section_title(self, doc: Document, title: str) -> None:
        """
        Agrega título de sección.

        Args:
            doc: Documento
            title: Título de la sección
        """
        heading = doc.add_heading(title, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def _add_parameters_table(self, doc: Document, parametros: Dict[str, Any]) -> None:
        """
        Agrega tabla de parámetros globales.

        Args:
            doc: Documento
            parametros: Dict con parámetros
        """
        # Definir campos y etiquetas
        fields = [
            ("nombre", "Nombre del Proyecto"),
            ("rubro", "Rubro / Sector"),
            ("ciudad", "Ciudad"),
            ("departamento", "Departamento"),
            ("pais", "País"),
            ("moneda", "Moneda"),
            ("tipo_cambio", "Tipo de Cambio (Bs/$us)"),
            ("inflacion", "Tasa de Inflación Anual"),
            ("horizonte", "Horizonte de Proyección (años)"),
            ("anio_base", "Año Base"),
            ("anio_inicio", "Año Inicio Operaciones"),
            ("impuesto_iue", "Impuesto IUE (%)"),
            ("impuesto_it", "Impuesto IT (%)"),
        ]

        # Crear tabla
        table = doc.add_table(rows=len(fields), cols=2)
        table.style = "Table Grid"

        # Llenar tabla
        for i, (field_key, field_label) in enumerate(fields):
            row = table.rows[i]

            # Etiqueta
            row.cells[0].text = field_label
            row.cells[0].paragraphs[0].runs[0].bold = True

            # Valor
            value = parametros.get(field_key)
            row.cells[1].text = self._format_value(value, field_key)

        # Ajustar anchos
        table.columns[0].width = Inches(2.5)
        table.columns[1].width = Inches(3.5)

    def _add_products_table(
        self, doc: Document, productos: List[Dict[str, Any]]
    ) -> None:
        """
        Agrega tabla de productos/servicios.

        Args:
            doc: Documento
            productos: Lista de productos
        """
        if not productos:
            doc.add_paragraph("No se registraron productos.")
            return

        # Crear tabla con encabezados
        table = doc.add_table(rows=len(productos) + 1, cols=4)
        table.style = "Table Grid"

        # Encabezados
        headers = [
            "N°",
            "Nombre del Producto",
            "Unidad de Medida",
            "Peso/Vol por Unidad",
        ]
        header_row = table.rows[0]

        for i, header in enumerate(headers):
            header_row.cells[i].text = header
            header_row.cells[i].paragraphs[0].runs[0].bold = True
            # Color de fondo
            self._set_cell_shading(header_row.cells[i], "D9E2F3")

        # Datos
        for i, producto in enumerate(productos):
            row = table.rows[i + 1]
            row.cells[0].text = str(producto.get("numero", i + 1))
            row.cells[1].text = str(producto.get("nombre", "") or "")
            row.cells[2].text = str(producto.get("unidad_medida", "Gramos") or "Gramos")
            row.cells[3].text = str(producto.get("peso_volumen", 125.0) or "125.0")

        # Ajustar anchos
        table.columns[0].width = Inches(0.5)
        table.columns[1].width = Inches(2.5)
        table.columns[2].width = Inches(1.5)
        table.columns[3].width = Inches(1.5)

    def _format_value(self, value: Any, field_key: str) -> str:
        """
        Formatea un valor para mostrar en Word.

        Args:
            value: Valor a formatear
            field_key: Clave del campo

        Returns:
            String formateado
        """
        if value is None:
            return ""

        # Porcentajes
        if field_key in ["inflacion", "impuesto_iue", "impuesto_it"]:
            if isinstance(value, (int, float)):
                return f"{float(value) * 100:.0f}%"
            return str(value)

        # Floats
        if field_key in ["tipo_cambio"]:
            if isinstance(value, (int, float)):
                return f"{float(value):.2f}"
            return str(value)

        # Strings y otros
        return str(value)

    def _set_cell_shading(self, cell, color: str) -> None:
        """
        Establece el color de fondo de una celda.

        Args:
            cell: Celda de la tabla
            color: Color en formato hex (ej: 'D9E2F3')
        """
        from docx.oxml.ns import qn
        from docx.oxml import parse_xml

        shading = parse_xml(
            f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'w:fill="{color}"/>'
        )
        cell._tc.get_or_add_tcPr().append(shading)

    def merge_from_template(
        self,
        template_path: Path,
        parametros: Dict[str, Any],
        productos: List[Dict[str, Any]],
        output_path: Path,
    ) -> Path:
        """
        Genera documento Word usando una plantilla existente.

        Args:
            template_path: Ruta a la plantilla .docx
            parametros: Dict con parámetros
            productos: Lista de productos
            output_path: Ruta donde guardar

        Returns:
            Path al documento generado

        Note:
            Por ahora, genera desde cero. En el futuro, puede usar
            python-docx-template para merge de plantillas.
        """
        # TODO: Implementar merge con plantilla
        # Por ahora, usar generate_from_excel_data
        return self.generate_from_excel_data(parametros, productos, output_path)
