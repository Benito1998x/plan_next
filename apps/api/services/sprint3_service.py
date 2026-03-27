"""
Sprint 3 Service — Indicadores de Mercado (Pipeline Estadístico Completo)

Orquesta el pipeline v2:
  1. Carga Bronze → Silver → Gold desde el survey
  2. Clasifica variables DINÁMICAMENTE via AIAgent (nominal/ordinal/Likert/cuantitativa)
  3. Aplica filtros de segmento si hay BuyerPersona en el plan
  4. Calcula indicadores: μ, σ, IC-95%, Z-score, distribuciones completas
  5. Genera Excel con hoja de indicadores
  6. Genera Word con sección estadística completa

El sistema es genérico: funciona con cualquier encuesta, no solo Shawarma.
"""

import json
import logging
from pathlib import Path
from typing import Dict, Any, Optional, List

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from sqlmodel import select

from database import get_database
from models.db.plan_data import Plan, DatosNegocio, BuyerPersona
from services.ai_agent import get_ai_agent
from services.pipeline.bronze import load_bronze_df_named
from services.pipeline.silver import build_silver
from services.pipeline.gold import build_gold, load_gold_variables
from services.pipeline.variable_classifier import classify_variables, load_cached_classification
from services.pipeline.segment_filter import filter_respondents, build_filters_from_buyer_persona
from services.pipeline.indicators_v2 import build_indicators_v2

logger = logging.getLogger(__name__)

# Ruta a plantilla 2.xlsx (base del Excel de tabulación con hoja "Muestra")
_PLANTILLAS_DIR = Path(__file__).parent.parent.parent.parent / "plantillas"
_PLANTILLA_2    = _PLANTILLAS_DIR / "excel" / "fase 2" / "plantilla 2.xlsx"

# Mapeo campo BuyerPersona → número de pregunta (heurístico genérico)
# Si la encuesta tiene una pregunta de edad → usualmente Q1; zona → Q4, etc.
# Se puede extender o hacer configurable por survey.
_BP_TO_Q_MAP = {
    "edad_objetivo":            1,   # Q1: edad
    "zona_residencia_objetivo": 4,   # Q4: zona de residencia
    "ocupacion_principal":      3,   # Q3: ocupación
}

# ─── Estilos Excel ────────────────────────────────────────────────────────────
_HEADER_FILL   = PatternFill(fill_type="solid", fgColor="FF1C246A")
_HEADER_FONT   = Font(bold=True, color="FFFFFFFF", size=10)
_SECTION_FILL  = PatternFill(fill_type="solid", fgColor="FF2F5496")
_SECTION_FONT  = Font(bold=True, color="FFFFFFFF", size=10)
_QUANT_FILL    = PatternFill(fill_type="solid", fgColor="FFE8F0FE")   # azul claro
_LIKERT_FILL   = PatternFill(fill_type="solid", fgColor="FFE2EFDA")   # verde claro
_ALT_FILL      = PatternFill(fill_type="solid", fgColor="FFF5F5F5")   # gris muy claro
_THIN_BORDER   = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

# ─── Colores Word ─────────────────────────────────────────────────────────────
_W_HEADER_BG = "1C246A"
_W_HEADER_FG = RGBColor(0xFF, 0xFF, 0xFF)
_W_QUANT_BG  = "DDEEFF"
_W_LIKERT_BG = "E2EFDA"
_W_ALT_BG    = "F5F5F5"


def run_sprint3(
    survey_id: int,
    plan_id: Optional[int],
    output_dir: Path,
    force_reclassify: bool = False,
) -> Dict[str, Any]:
    """
    Pipeline Sprint 3 completo.

    Args:
        survey_id         : ID del Survey en SQLite
        plan_id           : ID del Plan (para cargar BuyerPersona y filtrar)
        output_dir        : carpeta de salida
        force_reclassify  : si True, ignora el cache y re-clasifica con IA

    Returns:
        {
            "survey_id", "plan_id", "n_total", "n_segmento",
            "n_indicadores", "filtros_aplicados", "advertencias",
            "indicadores", "excel_path", "word_path"
        }
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    # ── 1. Bronze → Silver → Gold ─────────────────────────────────────────────
    logger.info(f"[Sprint3] Cargando Bronze para survey_id={survey_id}")
    df_raw         = load_bronze_df_named(survey_id)
    n_total        = len(df_raw)
    silver         = build_silver(df_raw)
    variable_names = load_gold_variables(survey_id)
    gold           = build_gold(silver, variable_names, survey_id=survey_id)

    # ── 2. Clasificación dinámica de variables (AIAgent) ──────────────────────
    agent = get_ai_agent()

    if not force_reclassify:
        classification = load_cached_classification(survey_id)

    if force_reclassify or not classification:
        logger.info("[Sprint3] Clasificando variables con AIAgent...")
        classification = classify_variables(silver, variable_names, survey_id, agent)

    # ── 3. Cargar plan y construir filtros de segmento ────────────────────────
    plan_nombre  = "Plan de Negocio"
    bp_ctx: Dict = {}
    dn_ctx: Dict = {}
    filtros      = {}
    advertencias: List[str] = []
    n_segmento   = None

    if plan_id is not None:
        db = get_database()
        with db.get_session() as session:
            plan = session.get(Plan, plan_id)
            if plan:
                plan_nombre = plan.nombre

            dn = session.exec(
                select(DatosNegocio).where(DatosNegocio.plan_id == plan_id)
            ).first()
            if dn:
                dn_ctx = {
                    "horario_atencion": dn.horario_atencion,
                    "zona_direccion":   dn.zona_direccion,
                    "canal_venta":      dn.canal_venta,
                    "capacidad_diaria": dn.capacidad_diaria,
                }

            bp = session.exec(
                select(BuyerPersona).where(BuyerPersona.plan_id == plan_id)
            ).first()
            if bp:
                bp_ctx = {
                    "edad_objetivo":            bp.edad_objetivo,
                    "genero_objetivo":          bp.genero_objetivo,
                    "ocupacion_principal":      bp.ocupacion_principal,
                    "zona_residencia_objetivo": bp.zona_residencia_objetivo,
                    "motivaciones_compra":      bp.motivaciones_compra,
                    "canal_informacion":        bp.canal_informacion,
                    "nivel_socioeconomico":     bp.nivel_socioeconomico,
                }

        # Construir filtros desde BuyerPersona
        if bp_ctx:
            filtros = build_filters_from_buyer_persona(bp_ctx, _BP_TO_Q_MAP)

    # ── 4. Segmentación (si hay filtros) ──────────────────────────────────────
    gold_segmento = gold   # por defecto, usar el gold completo

    if filtros:
        logger.info(f"[Sprint3] Aplicando filtros de segmento: {filtros}")
        try:
            df_filtrado, n_segmento, adv = filter_respondents(survey_id, filtros)
            advertencias.extend(adv)

            if n_segmento > 0:
                silver_seg  = build_silver(df_filtrado)
                # Para Gold segmentado: reusar los nombres de variables ya conocidos
                gold_segmento = build_gold(silver_seg, variable_names)
                logger.info(f"[Sprint3] Segmento: {n_segmento}/{n_total} encuestados")
            else:
                advertencias.append("El segmento no tiene encuestados. Se usará la muestra completa.")
                n_segmento = None
        except Exception as exc:
            logger.warning(f"[Sprint3] Error en filtrado: {exc}. Usando muestra completa.")
            advertencias.append(f"No se pudo filtrar por segmento: {exc}")

    # ── 5. Calcular indicadores v2 ────────────────────────────────────────────
    resultado = build_indicators_v2(
        gold=gold_segmento,
        classification=classification,
        n_total=n_total,
        n_segmento=n_segmento,
        filtros_aplicados=filtros if filtros else None,
        advertencias=advertencias,
    )

    # ── 6. Generar documentos ─────────────────────────────────────────────────
    slug       = plan_nombre.replace(" ", "_").lower()
    excel_path = output_dir / f"{slug}_indicadores_sprint3.xlsx"
    word_path  = output_dir / f"{slug}_indicadores_sprint3.docx"

    _generate_excel(resultado, plan_nombre, n_total, n_segmento, excel_path)
    _generate_word(resultado, plan_nombre, n_total, n_segmento, bp_ctx, dn_ctx, word_path)

    # ── 7. Guardar índice ─────────────────────────────────────────────────────
    index_path = output_dir / "index_sprint3.json"
    index = json.loads(index_path.read_text()) if index_path.exists() else {}
    index[str(survey_id)] = {"excel": str(excel_path), "word": str(word_path)}
    index_path.write_text(json.dumps(index, indent=2))

    return {
        "survey_id":       survey_id,
        "plan_id":         plan_id,
        "n_total":         n_total,
        "n_segmento":      n_segmento,
        "n_indicadores":   resultado["metadata"]["n_indicadores"],
        "filtros_aplicados": filtros if filtros else None,
        "advertencias":    advertencias,
        "indicadores":     resultado["indicadores"],
        "excel_path":      str(excel_path),
        "word_path":       str(word_path),
    }


# ─── Generación Excel ──────────────────────────────────────────────────────────

def _generate_excel(
    resultado: Dict,
    plan_nombre: str,
    n_total: int,
    n_segmento: Optional[int],
    output_path: Path,
) -> None:
    # Abrir plantilla 2.xlsx como base (tiene hojas "Tabulación" y "Muestra")
    # e insertar la hoja INDICADORES ANTES de "Muestra"
    if _PLANTILLA_2.exists():
        wb = openpyxl.load_workbook(str(_PLANTILLA_2))
        sheet_names = wb.sheetnames
        # Insertar INDICADORES en la posición de "Muestra" (desplaza "Muestra" a la derecha)
        muestra_idx = sheet_names.index("Muestra") if "Muestra" in sheet_names else len(sheet_names)
        ws = wb.create_sheet("INDICADORES", muestra_idx)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "INDICADORES"

    # Título
    ws.merge_cells("A1:F1")
    ws["A1"] = f"INDICADORES DE MERCADO — {plan_nombre.upper()}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 22

    base_str = f"Base total: N={n_total}"
    if n_segmento:
        base_str += f"  |  Segmento buyer persona: n={n_segmento}"
    ws.merge_cells("A2:F2")
    ws["A2"] = base_str
    ws["A2"].font = Font(italic=True, size=10)
    ws.row_dimensions[2].height = 16

    current_row = 4

    for ind in resultado["indicadores"]:
        tipo     = ind.get("tipo", "nominal")
        variable = ind.get("variable", "")
        q_num    = ind.get("pregunta", "")

        # ── Encabezado de variable ──────────────────────────────────────────
        tipo_label = {"nominal": "NOMINAL", "ordinal": "ORDINAL",
                      "ordinal_likert": "LIKERT", "cuantitativa": "CUANTITATIVA"}.get(tipo, tipo.upper())
        ws.merge_cells(f"A{current_row}:F{current_row}")
        ws[f"A{current_row}"] = f"  {q_num}. {variable}   [{tipo_label}]"
        ws[f"A{current_row}"].font  = _SECTION_FONT
        ws[f"A{current_row}"].fill  = _SECTION_FILL
        ws.row_dimensions[current_row].height = 16
        current_row += 1

        # ── Estadísticas cuantitativas / Likert ─────────────────────────────
        if tipo in ("cuantitativa", "ordinal_likert") and ind.get("mu") is not None:
            mu    = ind["mu"]
            sigma = ind.get("sigma", 0)
            unidad = ind.get("unidad", "")

            stats_fill = _QUANT_FILL if tipo == "cuantitativa" else _LIKERT_FILL

            row_labels = [
                ("Media ponderada (μ)", f"{mu:.4f} {unidad}"),
                ("Desv. estándar (σ)",  f"{sigma:.4f} {unidad}"),
                ("IC 95% inferior",     f"{ind.get('ic_95_lower', ''):.4f} {unidad}" if ind.get('ic_95_lower') is not None else "—"),
                ("IC 95% superior",     f"{ind.get('ic_95_upper', ''):.4f} {unidad}" if ind.get('ic_95_upper') is not None else "—"),
            ]
            if tipo == "ordinal_likert":
                row_labels += [
                    ("Score ponderado",  f"{ind.get('score_ponderado', mu):.2f} / {ind.get('max_score', 5)} pts"),
                    ("Score 0-100",      f"{ind.get('score_100', 0):.1f}%"),
                    ("TOP-2 Box",        f"{ind.get('top2_box', 0):.1f}%"),
                    ("BOTTOM-2 Box",     f"{ind.get('bottom2_box', 0):.1f}%"),
                ]

            for label, val in row_labels:
                ws[f"A{current_row}"] = label
                ws[f"A{current_row}"].font = Font(bold=True, size=9)
                ws[f"A{current_row}"].fill = stats_fill
                ws[f"A{current_row}"].border = _THIN_BORDER
                ws.merge_cells(f"B{current_row}:F{current_row}")
                ws[f"B{current_row}"] = val
                ws[f"B{current_row}"].font = Font(size=9)
                ws[f"B{current_row}"].border = _THIN_BORDER
                current_row += 1

        # ── Tabla de distribución ────────────────────────────────────────────
        # Encabezados de columnas
        col_headers = ["Opción", "Frec.", "%"]
        if tipo in ("cuantitativa", "ordinal_likert"):
            col_headers = ["Opción", "Valor num.", "Frec.", "%", "P(X≤val)", "P(X≥val)"]

        for col_i, hdr in enumerate(col_headers):
            cell = ws.cell(row=current_row, column=col_i + 1, value=hdr)
            cell.font   = _HEADER_FONT
            cell.fill   = _HEADER_FILL
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(horizontal="center")
        current_row += 1

        # Filas de distribución
        distribucion = ind.get("distribucion", [])
        for i, d in enumerate(distribucion):
            fill = _ALT_FILL if i % 2 == 0 else None
            row_data: List = [d["label"]]

            if tipo in ("cuantitativa", "ordinal_likert"):
                val_num = d.get("value")
                mu    = ind.get("mu")
                sigma = ind.get("sigma", 0)
                p_at_most  = ""
                p_at_least = ""
                if val_num is not None and mu is not None and sigma and sigma > 0:
                    from services.pipeline.quantitative_transform import cumulative_probability, survival_probability
                    p_at_most  = f"{cumulative_probability(val_num, mu, sigma)*100:.1f}%"
                    p_at_least = f"{survival_probability(val_num, mu, sigma)*100:.1f}%"
                row_data += [
                    f"{val_num:.2f}" if val_num is not None else "—",
                    d["freq"],
                    f"{d['pct']:.1f}%",
                    p_at_most,
                    p_at_least,
                ]
            else:
                row_data += [d["freq"], f"{d['pct']:.1f}%"]
                if tipo == "ordinal" and "pct_acumulado" in d:
                    row_data.append(f"{d['pct_acumulado']:.1f}%")

            for col_i, val in enumerate(row_data):
                cell = ws.cell(row=current_row, column=col_i + 1, value=val)
                cell.font   = Font(size=9)
                cell.border = _THIN_BORDER
                if fill:
                    cell.fill = fill

            current_row += 1

        # Moda
        ws[f"A{current_row}"] = f"→ Moda: {ind.get('moda', '')} ({ind.get('p_moda', 0):.1f}%)"
        ws[f"A{current_row}"].font = Font(italic=True, size=9)
        ws.merge_cells(f"A{current_row}:F{current_row}")

        # Nota
        if ind.get("nota"):
            current_row += 1
            ws.merge_cells(f"A{current_row}:F{current_row}")
            ws[f"A{current_row}"] = f"ℹ {ind['nota']}"
            ws[f"A{current_row}"].font = Font(italic=True, size=8, color="FF888888")

        current_row += 2   # Espacio entre indicadores

    # Ajustar anchos de columna
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 12
    ws.freeze_panes = "A4"

    wb.save(output_path)
    logger.info(f"[Sprint3] Excel generado: {output_path}")


# ─── Generación Word ───────────────────────────────────────────────────────────

def _generate_word(
    resultado: Dict,
    plan_nombre: str,
    n_total: int,
    n_segmento: Optional[int],
    buyer_persona: Dict,
    datos_negocio: Dict,
    output_path: Path,
) -> None:
    doc = Document()

    # Título
    title = doc.add_heading("INDICADORES DE MERCADO", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    base_str = f"Plan: {plan_nombre} | Base: N={n_total}"
    if n_segmento:
        base_str += f" | Segmento buyer persona: n={n_segmento}"
    sub = doc.add_paragraph(base_str)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True

    # Advertencias
    advertencias = resultado["metadata"].get("advertencias", [])
    if advertencias:
        doc.add_paragraph()
        for adv in advertencias:
            p = doc.add_paragraph()
            run = p.add_run(f"⚠ {adv}")
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xC0, 0x50, 0x00)

    doc.add_paragraph()
    doc.add_heading("Indicadores Detallados", level=1)

    for ind in resultado["indicadores"]:
        tipo     = ind.get("tipo", "nominal")
        variable = ind.get("variable", "")
        q_num    = ind.get("pregunta", "")
        n_preg   = ind.get("n_pregunta", "")

        tipo_label = {"nominal": "Nominal", "ordinal": "Ordinal",
                      "ordinal_likert": "Likert", "cuantitativa": "Cuantitativa"}.get(tipo, tipo)

        # Título de variable
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        run = p.add_run(f"{q_num}. {variable}  [{tipo_label}]  (n={n_preg})")
        run.bold = True
        run.font.size = Pt(11)

        # Interpretación
        if ind.get("interpretacion"):
            doc.add_paragraph(ind["interpretacion"]).runs[0].font.size = Pt(9)

        # Nota skip-logic
        if ind.get("nota"):
            np = doc.add_paragraph()
            nr = np.add_run(f"ℹ {ind['nota']}")
            nr.italic = True
            nr.font.size = Pt(8)

        # Estadísticas cuantitativas / Likert
        if tipo in ("cuantitativa", "ordinal_likert") and ind.get("mu") is not None:
            mu    = ind["mu"]
            sigma = ind.get("sigma", 0)
            unidad = ind.get("unidad", "")

            stats_lines = [
                f"μ (media) = {mu:.4f} {unidad}",
                f"σ (desv. estándar) = {sigma:.4f} {unidad}",
                f"IC 95%: [{ind.get('ic_95_lower', '—'):.4f}, {ind.get('ic_95_upper', '—'):.4f}] {unidad}",
            ]
            if tipo == "ordinal_likert":
                stats_lines += [
                    f"Score ponderado: {ind.get('score_ponderado', mu):.2f} pts",
                    f"Score 0-100: {ind.get('score_100', 0):.1f}%",
                    f"TOP-2 Box: {ind.get('top2_box', 0):.1f}%  |  BOTTOM-2 Box: {ind.get('bottom2_box', 0):.1f}%",
                ]

            bg = _W_QUANT_BG if tipo == "cuantitativa" else _W_LIKERT_BG
            stats_tbl = doc.add_table(rows=len(stats_lines), cols=1)
            stats_tbl.style = "Normal Table"
            for i, line in enumerate(stats_lines):
                cell = stats_tbl.rows[i].cells[0]
                cell.text = line
                cell.paragraphs[0].runs[0].font.size = Pt(9)
                _set_bg(cell, bg)
            stats_tbl.columns[0].width = Inches(6)
            doc.add_paragraph()

        # Tabla de distribución
        distribucion = ind.get("distribucion", [])
        if distribucion:
            has_val = tipo in ("cuantitativa", "ordinal_likert")
            has_acum = tipo == "ordinal"

            cols = 3 + (1 if has_val else 0) + (1 if has_acum else 0)
            tbl = doc.add_table(rows=len(distribucion) + 1, cols=cols)
            tbl.style = "Normal Table"
            _add_borders(tbl)

            # Encabezados
            headers = ["Opción"]
            if has_val:
                headers.append("Valor num.")
            headers += ["Frec.", "%"]
            if has_acum:
                headers.append("% Acum.")

            hrow = tbl.rows[0]
            widths = [Inches(2.4)] + [Inches(0.9)] * (cols - 1)
            for ci, (h, w) in enumerate(zip(headers, widths)):
                cell = hrow.cells[ci]
                cell.width = w
                _set_bg(cell, _W_HEADER_BG)
                p2 = cell.paragraphs[0]
                p2.clear()
                r = p2.add_run(h)
                r.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = _W_HEADER_FG
                p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Filas de datos
            for di, d in enumerate(distribucion):
                row = tbl.rows[di + 1]
                bg = _W_ALT_BG if di % 2 == 0 else None

                row_vals = [d["label"]]
                if has_val:
                    row_vals.append(f"{d.get('value', '—'):.2f}" if d.get("value") is not None else "—")
                row_vals += [str(d["freq"]), f"{d['pct']:.1f}%"]
                if has_acum:
                    row_vals.append(f"{d.get('pct_acumulado', 0):.1f}%")

                for ci, (val, w) in enumerate(zip(row_vals, widths)):
                    cell = row.cells[ci]
                    cell.width = w
                    if bg:
                        _set_bg(cell, bg)
                    p3 = cell.paragraphs[0]
                    p3.clear()
                    r = p3.add_run(str(val))
                    r.font.size = Pt(9)

        # Moda
        moda_para = doc.add_paragraph()
        moda_para.paragraph_format.space_before = Pt(4)
        mr = moda_para.add_run(f"→ Moda: {ind.get('moda', '')} ({ind.get('p_moda', 0):.1f}%)")
        mr.italic = True
        mr.font.size = Pt(9)

    # Pie de página
    doc.add_paragraph()
    fuente = doc.add_paragraph()
    fr = fuente.add_run(
        "Fuente: Encuesta de mercado procesada con pipeline Bronze→Silver→Gold. "
        "Estadísticas: Montgomery & Runger (2014) — Applied Statistics and Probability for Engineers."
    )
    fr.italic = True
    fr.font.size = Pt(8)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    logger.info(f"[Sprint3] Word generado: {output_path}")


# ── helpers Word ──────────────────────────────────────────────────────────────

def _add_borders(table) -> None:
    tbl   = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    for ex in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(ex)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def _set_bg(cell, color_hex: str) -> None:
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  color_hex)
    for ex in tcPr.findall(qn("w:shd")):
        tcPr.remove(ex)
    tcPr.append(shd)
