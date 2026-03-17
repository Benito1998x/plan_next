"""
Writer — Escribe las tablas de frecuencia Gold en un documento Word.

Carga el template Word (plan_15_generado.docx como base),
agrega una sección "3. Tabulación de la Encuesta" con un Cuadro por pregunta.

Formato de cada Cuadro:
  Cuadro N: Pregunta N — [Variable]
  ┌─────────────────────┬─────────────┬───────┐
  │ Detalle             │ Frecuencia  │   %   │
  ├─────────────────────┼─────────────┼───────┤
  │ 18 a 30 Años        │     130     │ 40.6% │
  │ 31 a 45 Años        │     110     │ 34.4% │
  │ ...                 │     ...     │  ...  │
  ├─────────────────────┼─────────────┼───────┤
  │ TOTAL               │     320     │100.0% │
  └─────────────────────┴─────────────┴───────┘
  Fuente: Encuesta de Google Form
"""

from pathlib import Path
from typing import Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# Colores (mismos que la plantilla Excel)
_HEADER_BG  = "1C246A"   # azul oscuro
_HEADER_FG  = RGBColor(0xFF, 0xFF, 0xFF)
_TOTAL_BG   = "D5D8DC"   # gris claro


def write_survey_to_word(
    template_path: Path,
    gold: Dict[int, Dict],
    output_path: Path,
) -> Path:
    """
    Carga el template Word, agrega la sección de tabulación y guarda.

    Args:
        template_path: plan_15_generado.docx (base del documento)
        gold: salida de gold.build_gold() — {q_num: {variable, options, total}}
        output_path: ruta del archivo de salida

    Returns:
        Path al archivo generado
    """
    doc = Document(template_path)

    # Separador entre el contenido existente y la tabulación
    doc.add_paragraph()
    heading = doc.add_heading("3. Tabulación de la Encuesta", level=1)
    heading.paragraph_format.space_before = Pt(18)

    for q_num in sorted(gold.keys()):
        _write_cuadro(doc, q_num, gold[q_num])

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def _write_cuadro(doc: Document, q_num: int, q_data: Dict) -> None:
    """Escribe un Cuadro N con su título, tabla y fuente."""
    variable = q_data["variable"]
    options  = q_data["options"]
    total    = q_data["total"]

    # ── Título ──────────────────────────────────────────────────────────────
    title_para = doc.add_paragraph()
    title_para.paragraph_format.space_before = Pt(14)
    title_para.paragraph_format.space_after  = Pt(2)
    run = title_para.add_run(f"Cuadro {q_num}: Pregunta {q_num} — {variable}")
    run.bold = True
    run.font.size = Pt(10)

    # ── Tabla: encabezado + opciones + total ─────────────────────────────────
    n_rows = 1 + len(options) + 1   # header + data + total
    table  = doc.add_table(rows=n_rows, cols=3)
    table.style = "Normal Table"
    _add_borders(table)

    # Anchos de columna
    col_widths = [Inches(3.2), Inches(1.4), Inches(0.9)]

    # Fila de encabezado
    header_row = table.rows[0]
    for col_i, (text, width) in enumerate(zip(["Detalle", "Frecuencia", "%"], col_widths)):
        cell = header_row.cells[col_i]
        cell.width = width
        _set_cell_bg(cell, _HEADER_BG)
        para = cell.paragraphs[0]
        para.clear()
        run = para.add_run(text)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = _HEADER_FG
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Filas de datos
    for i, opt in enumerate(options):
        data_row = table.rows[i + 1]
        _fill_data_row(data_row, opt["label"], opt["freq"], opt["pct"], col_widths)

    # Fila TOTAL
    total_row = table.rows[-1]
    _fill_data_row(total_row, "TOTAL", total, 1.0, col_widths, bold=True, bg=_TOTAL_BG)

    # ── Fuente ───────────────────────────────────────────────────────────────
    fuente = doc.add_paragraph()
    fuente.paragraph_format.space_before = Pt(2)
    fuente.paragraph_format.space_after  = Pt(6)
    run = fuente.add_run("Fuente: Encuesta de Google Form")
    run.italic = True
    run.font.size = Pt(9)


# ── helpers ────────────────────────────────────────────────────────────────

def _fill_data_row(row, label: str, freq: int, pct: float, widths, bold=False, bg=None):
    values = [label, str(freq), f"{pct * 100:.1f}%"]
    aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER]

    for col_i, (text, align, width) in enumerate(zip(values, aligns, widths)):
        cell = row.cells[col_i]
        cell.width = width
        if bg:
            _set_cell_bg(cell, bg)
        para = cell.paragraphs[0]
        para.clear()
        run = para.add_run(text)
        run.font.size = Pt(10)
        if bold:
            run.bold = True
        para.alignment = align


def _set_cell_bg(cell, color_hex: str) -> None:
    """Aplica color de fondo a una celda via XML."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  color_hex)
    # Reemplaza shd existente si hay
    for existing in tcPr.findall(qn("w:shd")):
        tcPr.remove(existing)
    tcPr.append(shd)


def _add_borders(table) -> None:
    """Agrega bordes visibles a toda la tabla via XML."""
    tbl  = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)

    for existing in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(existing)

    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"),   "single")
        el.set(qn("w:sz"),    "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
