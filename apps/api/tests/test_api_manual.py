"""
Test manual de la API sin frontend.

Usa TestClient de FastAPI/Starlette para llamar los endpoints
directamente, sin levantar un servidor HTTP.

Prueba el flujo completo:
  POST /api/v1/process-plan  →  genera Excel + Word
  GET  /api/v1/download/{id}/excel
  GET  /api/v1/download/{id}/word

Ejecutar desde apps/api/:
  python tests/test_api_manual.py
"""

import sys
import os
import json
from pathlib import Path

# Setup path para imports
sys.path.insert(0, str(Path(__file__).parent.parent))

# Cargar .env del root del proyecto
from dotenv import load_dotenv, find_dotenv
load_dotenv(find_dotenv(usecwd=True) or find_dotenv())

from fastapi.testclient import TestClient
from main import app

client = TestClient(app)

# ─── Carpeta de output para los archivos generados ───────────────────────────
OUTPUT_DIR = Path(__file__).parent.parent.parent.parent / "plantillas" / "excel"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def separator(title: str):
    print(f"\n{'='*55}")
    print(f"  {title}")
    print('='*55)


def test_health():
    separator("1. Health Check")
    r = client.get("/health")
    print(f"  Status: {r.status_code}")
    print(f"  Body:   {r.json()}")
    assert r.status_code == 200, "Health check failed"
    print("  ✓ OK")


def test_process_plan_basico():
    """Caso base: datos mínimos, agente completa el resto."""
    separator("2. POST /process-plan — datos mínimos")

    payload = {
        "nombre": "Shawarma Cruz",
        "rubro": "Gastronomía",
        "num_productos": 2,
    }
    print(f"  Input:  {json.dumps(payload, ensure_ascii=False)}")

    r = client.post("/api/v1/process-plan", json=payload)
    print(f"  Status: {r.status_code}")

    assert r.status_code == 200, f"Error: {r.text}"
    data = r.json()["data"]

    print(f"  plan_id:    {data['plan_id']}")
    print(f"  nombre:     {data['nombre']}")
    print(f"  rubro:      {data['rubro']}")
    print(f"  ciudad:     {data['parametros']['ciudad']}")
    print(f"  productos:  {[p['nombre'] for p in data['productos']]}")

    if data.get("_agent_error"):
        print(f"  ⚠ Agent fallback activo: {data['_agent_error'][:80]}")
    else:
        print(f"  ✓ Agente IA respondió OK (modelo: {os.getenv('OPENAI_MODEL', 'gpt-4o-mini')})")

    print("  ✓ proceso_plan OK")
    return data["plan_id"]


def test_process_plan_con_productos():
    """Caso con productos detallados (sin inferencia del agente)."""
    separator("3. POST /process-plan — productos detallados")

    payload = {
        "nombre": "Shawarma Cruz",
        "rubro": "Gastronomía",
        "ciudad": "Santa Cruz de la Sierra",
        "productos": [
            {"nombre": "Shawarma De Carne De Cordero", "unidad_medida": "Gramos", "peso_volumen": 250},
            {"nombre": "Shawarma De Carne De Pollo",   "unidad_medida": "Gramos", "peso_volumen": 250},
            {"nombre": "Shawarma De Carne De Res",     "unidad_medida": "Gramos", "peso_volumen": 250},
        ],
    }
    print(f"  Productos enviados: {len(payload['productos'])}")

    r = client.post("/api/v1/process-plan", json=payload)
    assert r.status_code == 200, f"Error: {r.text}"

    data = r.json()["data"]
    print(f"  Productos recibidos: {[p['nombre'] for p in data['productos']]}")
    print("  ✓ OK")
    return data["plan_id"]


def test_download_excel(plan_id: int):
    """Descarga el Excel generado y verifica su contenido."""
    separator(f"4. GET /download/{plan_id}/excel")

    r = client.get(f"/api/v1/download/{plan_id}/excel")
    print(f"  Status: {r.status_code}")
    assert r.status_code == 200, f"Error: {r.text}"
    assert "spreadsheetml" in r.headers["content-type"]

    # Guardar en plantillas/excel para inspección visual
    out_path = OUTPUT_DIR / f"plan_{plan_id}_generado.xlsx"
    out_path.write_bytes(r.content)
    print(f"  Guardado en: {out_path}")

    # Verificar contenido
    from openpyxl import load_workbook
    wb = load_workbook(out_path, data_only=True)
    ws = wb["INICIO"]

    nombre = ws["B4"].value
    rubro  = ws["B5"].value
    prod1  = ws["B20"].value
    prod2  = ws["B21"].value

    print(f"  B4  (nombre): {nombre}")
    print(f"  B5  (rubro):  {rubro}")
    print(f"  B20 (prod 1): {prod1}")
    print(f"  B21 (prod 2): {prod2}")
    print(f"  B10 (tipo_cambio): {ws['B10'].value}")
    print(f"  B15 (IUE): {ws['B15'].value}")

    assert nombre is not None, "B4 (nombre) está vacío"
    assert rubro  is not None, "B5 (rubro) está vacío"
    assert prod1  is not None, "B20 (prod1) está vacío"
    wb.close()
    print("  ✓ Contenido Excel verificado")


def test_download_word(plan_id: int):
    """Descarga el Word generado y verifica su contenido."""
    separator(f"5. GET /download/{plan_id}/word")

    r = client.get(f"/api/v1/download/{plan_id}/word")
    print(f"  Status: {r.status_code}")
    assert r.status_code == 200, f"Error: {r.text}"
    assert "wordprocessingml" in r.headers["content-type"]

    # Guardar junto al Excel para inspección
    word_dir = OUTPUT_DIR.parent / "word"
    word_dir.mkdir(exist_ok=True)
    out_path = word_dir / f"plan_{plan_id}_generado.docx"
    out_path.write_bytes(r.content)
    print(f"  Guardado en: {out_path}")

    # Verificar contenido
    from docx import Document
    doc = Document(out_path)

    print(f"  Párrafos con texto:")
    for p in doc.paragraphs:
        if p.text.strip():
            print(f"    [{p.style.name}] {p.text[:60]}")

    print(f"  Tablas: {len(doc.tables)}")
    for i, t in enumerate(doc.tables):
        print(f"    Tabla {i}: {len(t.rows)}r × {len(t.columns)}c")
        for row in t.rows[:4]:
            print(f"      {[c.text[:20] for c in row.cells]}")

    assert len(doc.tables) >= 1, "El Word no tiene tablas"
    print("  ✓ Contenido Word verificado")


def test_upload_rellenado():
    """
    Prueba con la plantilla de ejemplo ya rellenada
    (plantilla 1 -rellenado.xlsx) como si fuera el cliente.
    """
    separator("6. POST /upload-excel — con plantilla rellenada de ejemplo")

    rellenado = Path(__file__).parent.parent.parent.parent / "plantillas" / "excel" / "plantilla 1 -rellenado.xlsx"
    if not rellenado.exists():
        print(f"  ⚠ No existe {rellenado}, skipping")
        return

    with open(rellenado, "rb") as f:
        r = client.post(
            "/api/v1/upload-excel",
            files={"file": ("plantilla 1 -rellenado.xlsx", f,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    print(f"  Status: {r.status_code}")
    if r.status_code == 200:
        data = r.json()["data"]
        print(f"  plan_id:   {data.get('plan_id')}")
        print(f"  nombre:    {data.get('nombre')}")
        print(f"  productos: {[p['nombre'] for p in data.get('productos', [])]}")
        print("  ✓ OK")
    else:
        print(f"  ✗ Error: {r.text[:200]}")


if __name__ == "__main__":
    print("\n=== Test manual BPAE API - sin frontend ===")
    print(f"   Modelo IA: {os.getenv('OPENAI_MODEL', 'gpt-4o-mini')}")
    print(f"   Provider:  {os.getenv('LLM_PROVIDER', 'openai')}")

    try:
        test_health()
        plan_id_basico = test_process_plan_basico()
        plan_id_detallado = test_process_plan_con_productos()
        test_download_excel(plan_id_detallado)
        test_download_word(plan_id_detallado)
        test_upload_rellenado()

        separator("RESULTADO FINAL")
        print("  ✓ Todos los tests pasaron")
        print(f"  Archivos generados en: {OUTPUT_DIR.parent}")

    except AssertionError as e:
        print(f"\n  ✗ FALLO: {e}")
        sys.exit(1)
    except Exception as e:
        print(f"\n  ✗ ERROR INESPERADO: {e}")
        import traceback; traceback.print_exc()
        sys.exit(1)
