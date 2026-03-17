"""
Test Sprint 2 — Pipeline Bronze → Silver → Gold → Word

Valida:
  1. Bronze   : 320 filas cargadas en SQLite correctamente
  2. Silver   : 16 preguntas, frecuencias correctas, limpieza de prefijos
  3. Gold     : variables IA inferidas para las 16 preguntas
  4. Writer   : Word generado con 16 Cuadros, estructura correcta
  5. API      : POST /process-survey → word_url funcional
  6. Download : GET /download/survey/{id}/word → .docx válido

Ejecutar desde apps/api/:
  ../../venv/Scripts/python -X utf8 tests/test_survey.py
"""

import sys
import sqlite3
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from fastapi.testclient import TestClient
from main import app

client = TestClient(app)

_PLANTILLAS   = Path(__file__).parent.parent.parent.parent / "plantillas"
ENCUESTA_PATH = _PLANTILLAS / "excel" / "fase 2" / "Encuesta.xlsx"
WORD_TEMPLATE = _PLANTILLAS / "word"  / "fase 2" / "plan_15_generado.docx"
OUTPUT_DIR    = _PLANTILLAS / "word"  / "fase 2"
DB_PATH       = Path(__file__).parent.parent / "data" / "plan.db"

EXPECTED_ROWS      = 320
EXPECTED_QUESTIONS = 16


def sep(title: str):
    print(f"\n{'='*60}")
    print(f"  {title}")
    print("=" * 60)


def ok(msg: str = ""):
    print(f"  OK {msg}")


# ─────────────────────────────────────────────────────────────────────────────
# 1. BRONZE — Ingesta en SQLite
# ─────────────────────────────────────────────────────────────────────────────

def test_bronze():
    sep("1. BRONZE — Ingesta Excel → SQLite")

    assert ENCUESTA_PATH.exists(), f"No existe: {ENCUESTA_PATH}"

    from services.pipeline.bronze import ingest_excel, extract_question_texts
    import pandas as pd

    # Leer raw para comparar
    df_raw = pd.read_excel(ENCUESTA_PATH, engine="openpyxl").dropna(how="all")
    print(f"  Filas en Excel    : {len(df_raw)}")

    # Ingestar en SQLite
    survey_id, df_ingested = ingest_excel(ENCUESTA_PATH, "Test Sprint 2")
    print(f"  survey_id SQLite  : {survey_id}")
    print(f"  Filas ingresadas  : {len(df_ingested)}")

    assert len(df_ingested) == EXPECTED_ROWS, (
        f"Esperaba {EXPECTED_ROWS} filas, obtuvo {len(df_ingested)}"
    )

    # Verificar en SQLite que se guardaron las respuestas
    # (sólo se guardan respuestas no-nulas, así que puede ser < 320×16)
    import re as _re
    q_cols = [c for c in df_ingested.columns if _re.match(r'^\d+\.', str(c))]
    expected_count = int(sum(df_ingested[c].notna().sum() for c in q_cols))

    con = sqlite3.connect(DB_PATH)
    (count,) = con.execute(
        "SELECT COUNT(*) FROM surveyresponse WHERE survey_id=?", (survey_id,)
    ).fetchone()
    con.close()
    print(f"  Respuestas en DB  : {count}  (no-nulas esperadas: {expected_count})")
    assert count == expected_count, (
        f"Esperaba {expected_count} filas no-nulas en surveyresponse, obtuvo {count}"
    )

    # Verificar textos de preguntas
    texts = extract_question_texts(df_ingested)
    print(f"  Preguntas detectadas: {len(texts)}")
    assert len(texts) == EXPECTED_QUESTIONS

    ok(f"— survey_id={survey_id}, {count} respuestas en SQLite")
    return survey_id, df_ingested


# ─────────────────────────────────────────────────────────────────────────────
# 2. SILVER — Limpieza + Frecuencias
# ─────────────────────────────────────────────────────────────────────────────

def test_silver(df_raw):
    sep("2. SILVER — Limpieza pandas + groupby")

    from services.pipeline.silver import clean_label, build_silver

    # Verificar clean_label
    casos = [
        ("a) 18 a 30 Años",         "18 a 30 Años"),
        ("b) Masculino",             "Masculino"),
        ("c) Shawarma De Cordero",   "Shawarma De Cordero"),
        ("Sin prefijo",              "Sin prefijo"),
    ]
    for raw, expected in casos:
        resultado = clean_label(raw)
        assert resultado == expected, f"clean_label('{raw}') = '{resultado}' != '{expected}'"
    print("  clean_label: 4/4 casos correctos")

    # Construir silver completo
    silver = build_silver(df_raw)

    print(f"  total_responses   : {silver['total_responses']}")
    print(f"  preguntas silver  : {len(silver['questions'])}")

    assert silver["total_responses"] == EXPECTED_ROWS
    assert len(silver["questions"]) == EXPECTED_QUESTIONS

    # Verificar P1 (Edad) — frecuencias deben sumar 320 y porcentajes 100%
    p1 = silver["questions"][1]
    total_p1 = sum(opt["frecuencia"] for _, opt in p1["df_freq"].iterrows())
    total_pct = sum(opt["porcentaje"] for _, opt in p1["df_freq"].iterrows())

    print(f"\n  P1 '{p1['texto_pregunta'][:40]}':")
    print(f"    Opciones         : {len(p1['df_freq'])}")
    print(f"    Suma frecuencias : {total_p1}  (esperado: {EXPECTED_ROWS})")
    print(f"    Suma porcentajes : {total_pct:.4f}  (esperado: 1.0)")

    for _, row in p1["df_freq"].iterrows():
        print(f"    {str(row['label'])[:35]:35s} | {int(row['frecuencia']):4d} | {row['porcentaje']*100:5.1f}%")

    assert total_p1 == EXPECTED_ROWS, f"Suma frecuencias P1 = {total_p1}"
    assert abs(total_pct - 1.0) < 0.001, f"Suma porcentajes P1 = {total_pct}"

    # Verificar que no hay prefijos en los labels
    for q_num, q_data in silver["questions"].items():
        for _, row in q_data["df_freq"].iterrows():
            label = str(row["label"])
            assert not (len(label) >= 2 and label[1] == ")"), (
                f"P{q_num}: label con prefijo no limpiado: '{label}'"
            )

    print("\n  Sin prefijos 'a) ' en ninguna label: OK")
    ok(f"— {EXPECTED_QUESTIONS} preguntas, frecuencias correctas")
    return silver


# ─────────────────────────────────────────────────────────────────────────────
# 3. GOLD — Variables IA + Cache SQLite
# ─────────────────────────────────────────────────────────────────────────────

def test_gold(silver, survey_id):
    sep("3. GOLD — Variables IA + Cache SQLite")

    from services.ai_agent import get_ai_agent
    from services.pipeline.bronze import extract_question_texts
    from services.pipeline.gold import build_gold, load_gold_variables
    import pandas as pd

    # Inferir variables
    q_texts = {
        q_num: q["texto_pregunta"]
        for q_num, q in silver["questions"].items()
    }
    agent = get_ai_agent()
    variable_names = agent.infer_variable_names(q_texts)

    print(f"  Variables inferidas: {len(variable_names)}")
    assert len(variable_names) == EXPECTED_QUESTIONS

    for num, var in sorted(variable_names.items()):
        print(f"    P{num:>2}: {var}")

    # Construir Gold
    gold = build_gold(silver, variable_names, survey_id=survey_id)

    # Verificar cache en SQLite
    cached = load_gold_variables(survey_id)
    print(f"\n  Variables en cache SQLite: {len(cached)}")
    assert len(cached) == EXPECTED_QUESTIONS, "Cache SQLite incompleto"

    # Verificar estructura Gold por pregunta
    for q_num, q_data in gold.items():
        assert "variable" in q_data
        assert "options" in q_data
        assert "total" in q_data
        assert q_data["total"] > 0
        # Suma de frecuencias == total
        suma = sum(o["freq"] for o in q_data["options"])
        assert suma == q_data["total"], (
            f"P{q_num}: suma frecuencias {suma} != total {q_data['total']}"
        )
        # Porcentajes suman 1.0
        suma_pct = sum(o["pct"] for o in q_data["options"])
        assert abs(suma_pct - 1.0) < 0.001, (
            f"P{q_num}: suma porcentajes {suma_pct:.4f} != 1.0"
        )

    print("  Estructura Gold: todas las preguntas OK")
    ok(f"— Gold correcto, cache en SQLite")
    return gold


# ─────────────────────────────────────────────────────────────────────────────
# 4. WRITER — Word con 16 Cuadros
# ─────────────────────────────────────────────────────────────────────────────

def test_writer(gold):
    sep("4. WRITER — Documento Word")

    from services.pipeline.writer import write_survey_to_word
    from docx import Document

    assert WORD_TEMPLATE.exists(), f"Template no existe: {WORD_TEMPLATE}"

    out_path = OUTPUT_DIR / "test_sprint2_output.docx"
    result_path = write_survey_to_word(WORD_TEMPLATE, gold, out_path)

    assert result_path.exists(), "El archivo Word no fue creado"
    size_kb = result_path.stat().st_size / 1024
    print(f"  Archivo generado  : {result_path.name}")
    print(f"  Tamaño            : {size_kb:.1f} KB")
    assert size_kb > 10, "El archivo parece vacío"

    # Verificar estructura del Word
    doc = Document(result_path)

    # Contar tablas: template ya tiene 2 (datos globales + productos de Sprint 1)
    # Writer agrega 16 más → total esperado = 18
    from docx import Document as _Doc
    template_tables = len(_Doc(WORD_TEMPLATE).tables)
    n_tables = len(doc.tables)
    survey_tables = n_tables - template_tables
    print(f"  Tablas template   : {template_tables}")
    print(f"  Tablas agregadas  : {survey_tables}  (esperado: {EXPECTED_QUESTIONS})")
    print(f"  Total tablas doc  : {n_tables}")
    assert survey_tables == EXPECTED_QUESTIONS, (
        f"Esperaba {EXPECTED_QUESTIONS} tablas de encuesta, se agregaron {survey_tables}"
    )

    # Verificar primera tabla de encuesta (template tiene 'template_tables' tablas antes)
    t0 = doc.tables[template_tables]   # primera tabla de la sección de encuesta
    header_row = t0.rows[0]
    headers = [c.text.strip() for c in header_row.cells]
    print(f"  Headers tabla P1  : {headers}")
    assert headers == ["Detalle", "Frecuencia", "%"], f"Headers incorrectos: {headers}"

    # Verificar fila TOTAL
    last_row = t0.rows[-1]
    total_label = last_row.cells[0].text.strip()
    assert total_label == "TOTAL", f"Última fila no es TOTAL: '{total_label}'"
    print(f"  Fila TOTAL        : {[c.text.strip() for c in last_row.cells]}")

    # Verificar que hay 16 Cuadros en los párrafos
    cuadro_paras = [
        p.text for p in doc.paragraphs
        if p.text.startswith("Cuadro ")
    ]
    print(f"  Titulos Cuadro N  : {len(cuadro_paras)}")
    assert len(cuadro_paras) == EXPECTED_QUESTIONS
    print(f"  Primer cuadro     : {cuadro_paras[0]}")
    print(f"  Ultimo cuadro     : {cuadro_paras[-1]}")

    ok(f"— Word con {n_tables} tablas, estructura correcta")
    return result_path


# ─────────────────────────────────────────────────────────────────────────────
# 5. API — POST /process-survey
# ─────────────────────────────────────────────────────────────────────────────

def test_api_process_survey():
    sep("5. API — POST /process-survey")

    assert ENCUESTA_PATH.exists()

    with open(ENCUESTA_PATH, "rb") as f:
        r = client.post(
            "/api/v1/process-survey",
            files={"file": (
                "Encuesta.xlsx", f,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )},
        )

    print(f"  Status            : {r.status_code}")
    if r.status_code != 200:
        print(f"  Error             : {r.text[:400]}")
    assert r.status_code == 200

    data = r.json()["data"]
    print(f"  survey_id         : {data['survey_id']}")
    print(f"  bronze_rows       : {data['bronze_rows']}")
    print(f"  silver_questions  : {data['silver_questions']}")
    print(f"  word_url          : {data.get('word_url')}")

    assert data["bronze_rows"] == EXPECTED_ROWS
    assert data["silver_questions"] == EXPECTED_QUESTIONS
    assert "word_url" in data, "Falta word_url en la respuesta"

    ok(f"— survey_id={data['survey_id']}, word_url presente")
    return data["survey_id"]


# ─────────────────────────────────────────────────────────────────────────────
# 6. DOWNLOAD — GET /download/survey/{id}/word
# ─────────────────────────────────────────────────────────────────────────────

def test_api_download(survey_id: int):
    sep(f"6. API — GET /download/survey/{survey_id}/word")

    r = client.get(f"/api/v1/download/survey/{survey_id}/word")
    print(f"  Status            : {r.status_code}")
    assert r.status_code == 200, f"Error: {r.text}"
    assert "wordprocessingml" in r.headers.get("content-type", ""), (
        f"Content-Type incorrecto: {r.headers.get('content-type')}"
    )

    out_path = OUTPUT_DIR / f"api_tabulacion_{survey_id}.docx"
    out_path.write_bytes(r.content)
    size_kb = len(r.content) / 1024
    print(f"  Archivo descargado: {out_path.name}")
    print(f"  Tamaño            : {size_kb:.1f} KB")
    assert size_kb > 10, "Archivo descargado parece vacío"

    ok(f"— .docx descargado ({size_kb:.1f} KB)")


# ─────────────────────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    print("\n" + "="*60)
    print("  TEST SPRINT 2 — Pipeline Bronze→Silver→Gold→Word")
    print("="*60)
    print(f"  Encuesta : {ENCUESTA_PATH}")
    print(f"  Template : {WORD_TEMPLATE}")
    print(f"  Output   : {OUTPUT_DIR}")
    print(f"  SQLite   : {DB_PATH}")

    failed = []

    try:
        # Tests de capas (directos, sin HTTP)
        survey_id, df_raw = test_bronze()
        silver = test_silver(df_raw)
        gold   = test_gold(silver, survey_id)
        test_writer(gold)

        # Tests de API
        api_survey_id = test_api_process_survey()
        test_api_download(api_survey_id)

    except AssertionError as e:
        failed.append(str(e))
        print(f"\n  ASSERTION FALLIDA: {e}")
        import traceback; traceback.print_exc()
    except Exception as e:
        failed.append(str(e))
        print(f"\n  ERROR: {e}")
        import traceback; traceback.print_exc()

    sep("RESULTADO FINAL")
    if failed:
        print(f"  FALLARON {len(failed)} test(s):")
        for f in failed:
            print(f"  - {f}")
        sys.exit(1)
    else:
        print("  Todos los tests pasaron correctamente")
        print(f"  Archivos generados en: {OUTPUT_DIR}")
